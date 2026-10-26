"""Bounded parser subprocess; receives a private temporary filename, emits JSON."""
import json
from pathlib import Path
import sys
import zipfile


def parse(path):
    suffix = path.suffix.lower()
    if suffix in ('.md', '.txt'):
        pages = [path.read_text(encoding='utf-8')]
    elif suffix == '.ipynb':
        import nbformat
        nb = nbformat.read(path, as_version=4)
        nbformat.validate(nb)
        if len(nb.cells) > 100:
            raise ValueError('too many notebook cells')
        return {'pages': [c.source for c in nb.cells], 'notebook': json.loads(nbformat.writes(nb))}
    elif suffix == '.pdf':
        if not path.read_bytes().startswith(b'%PDF-'):
            raise ValueError('invalid PDF signature')
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > 200:
                raise ValueError('too many pages')
            pages = [p.extract_text() or '' for p in pdf.pages]
    elif suffix == '.pptx':
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50_000_000:
                raise ValueError('expanded file too large')
        from pptx import Presentation
        prs = Presentation(str(path))
        if len(prs.slides) > 200:
            raise ValueError('too many slides')
        pages = ['\n'.join(s.text for s in slide.shapes if s.has_text_frame) for slide in prs.slides]
    else:
        raise ValueError('unsupported file')
    if not any(p.strip() for p in pages):
        raise ValueError('no extractable text; OCR required')
    if sum(map(len, pages)) > 500_000:
        raise ValueError('extracted text too large')
    return {'pages': pages}


if __name__ == '__main__':
    try:
        print(json.dumps(parse(Path(sys.argv[1])), ensure_ascii=False))
    except Exception:
        print('Document cannot be parsed or exceeds limits', file=sys.stderr)
        sys.exit(1)
